import sys
from pathlib import Path

import openpyxl
import pytest

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from model import engine  # noqa: E402
from model.inputs import DEFAULT_WORKBOOK, InputState, load_workbook_data  # noqa: E402


@pytest.fixture(scope="session")
def wb():
    return load_workbook_data(DEFAULT_WORKBOOK)


@pytest.fixture(scope="session")
def formulas():
    return openpyxl.load_workbook(DEFAULT_WORKBOOK)


@pytest.fixture(scope="session")
def cached():
    return openpyxl.load_workbook(DEFAULT_WORKBOOK, data_only=True)


@pytest.fixture()
def state(wb):
    return InputState.from_specs(wb.inputs)


@pytest.fixture(scope="session")
def base(wb):
    """Engine result at the workbook's active values and default package."""
    values = InputState.from_specs(wb.inputs).values()
    return engine.run_model(values, wb.include, wb.units, wb.curves)
