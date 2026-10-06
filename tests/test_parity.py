"""Parity: the engine must give the workbook's numbers for the workbook's inputs."""
import shutil
import subprocess
import time

import openpyxl
import pytest

from model import engine, sensitivity
from model.inputs import DEFAULT_WORKBOOK, InputState
from workbook_map import TEXT_ONLY_FORMULAS, expectations

MODEL_SHEETS = ["Inputs", "Hazards", "Value chain", "Calc", "Supply", "Engine", "Portfolio", "Results", "Checks"]

# Tighter than the brief asks for (USD 1, or 1e-6 for ratios): the engine is a direct translation.
ABS_TOL = 1e-6
REL_TOL = 1e-9


def _mismatches(result, state, wb, sheets):
    """Compare every mapped cell with the given workbook (values only). Returns (count, problems)."""
    problems = []
    count = 0
    for sheet, cell, got in expectations(result, state, wb.inputs, result["portfolio"]["include"]):
        want = sheets[sheet][cell].value
        count += 1
        if isinstance(got, str) or isinstance(want, str):
            if got != want:
                problems.append(f"{sheet}!{cell}: workbook {want!r}, engine {got!r}")
        elif want is None or abs(got - want) > ABS_TOL + REL_TOL * abs(want):
            problems.append(f"{sheet}!{cell}: workbook {want!r}, engine {got!r}")
    return count, problems


# --- case 1: base case against Excel's cached values --------------------------------------

def test_base_case_matches_every_cached_cell(base, state, wb, cached):
    count, problems = _mismatches(base, state, wb, cached)
    assert not problems, f"{len(problems)} of {count} cells differ:\n" + "\n".join(problems[:25])
    assert count > 6000


def test_every_formula_cell_is_covered(base, state, wb, formulas):
    """No formula in the model sheets is left out of the parity map."""
    mapped = {(sheet, cell) for sheet, cell, _ in expectations(base, state, wb.inputs, wb.include)}
    missing = []
    for sheet in MODEL_SHEETS:
        for row in formulas[sheet].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    key = (sheet, c.coordinate)
                    if key not in mapped and key not in TEXT_ONLY_FORMULAS:
                        missing.append(f"{sheet}!{c.coordinate}")
    assert not missing, f"{len(missing)} formula cells are not covered: {missing[:20]}"


EXPECTED_NPV = {  # from the brief: base case, default package
    "1": -135_111.79, "2": -31_252.92, "3": -45_446.69, "4": -126_156.42, "5": -43_059.93, "6": -66_504.12,
    "7a": 72_102.62, "7b": 22_445.39, "7c": 37_002.52, "7d": 139_263.34, "8": 123_164.62, "9": -98_388.99,
    "10": 1_427.05,
}


def test_headline_values_from_the_brief(base, wb):
    assert {oid for oid, on in wb.include.items() if on} == {"2", "5", "7a", "7d", "8", "10"}
    for oid, want in EXPECTED_NPV.items():
        assert base["options"][oid]["npv"] == pytest.approx(want, abs=0.01), f"option {oid}"
    assert base["portfolio"]["npv"] == pytest.approx(260_114.90, abs=0.01)
    assert base["portfolio"]["npv_with_suppliers"] == pytest.approx(363_769.33, abs=0.01)
    calc = base["calc"]
    assert calc["c_margin"] == pytest.approx(5_215.45, abs=0.01)
    assert calc["c_VST"] == pytest.approx(2_037.14, abs=0.01)
    assert calc["c_out_cost"] == pytest.approx(33_910.16, abs=0.01)
    assert calc["c_floss0"] == pytest.approx(123_290.91, abs=0.01)
    share = base["value_chain"]["summary"]["share_gh"]
    assert share["weighted"] == pytest.approx(0.30993, abs=1e-5)
    assert share["amazon"] == pytest.approx(0.269, abs=5e-4)
    assert share["domestic"] == pytest.approx(0.895, abs=5e-4)
    assert base["value_chain"]["business"]["share_stays"] == pytest.approx(0.400, abs=5e-4)


def test_calc_names_point_at_the_mapped_cells(base, cached):
    """Each Calc/Value chain defined name resolves to a cell whose cached value the engine reproduces."""
    vc = base["value_chain"]
    by_name = dict(base["calc"])
    by_name.update({
        "c_crude_kg": vc["crude_kg"], "c_units": vc["units"], "c_units_amz": vc["units_amz"], "c_fob": vc["fob"],
        "c_gh_share_exp": vc["amazon"]["share_gh"], "c_gh_share_dom": vc["domestic"]["share_gh"],
        "c_unit_margin": vc["summary"]["brand_margin"]["weighted"], "c_margin_vc": vc["margin_vc"],
        "c_sales": vc["business"]["total_revenue"],
        "p_npv": base["portfolio"]["npv"], "p_npv_cons": base["portfolio"]["npv_with_suppliers"],
    })
    names = [n for n in cached.defined_names if not n.startswith("in_")]
    assert sorted(names) == sorted(by_name)
    for name in names:
        sheet, coord = next(iter(cached.defined_names[name].destinations))
        want = cached[sheet][coord.replace("$", "")].value
        assert by_name[name] == pytest.approx(want, rel=1e-9, abs=1e-6), name


def test_constants_match_the_workbook(wb):
    assert wb.curves == engine.JRC_CURVES
    assert wb.option_names == engine.OPTION_NAMES
    assert len(wb.inputs) == 189
    assert [s.id for s in wb.inputs if s.follows_scenario] == ["H01", "H02", "H04", "H05", "H06", "H08"]


# --- sensitivity snapshot (static values in the workbook) ---------------------------------

def test_scenario_snapshot(state, wb):
    runs = sensitivity.scenario_runs(state, wb.include, wb.units, wb.curves)
    rows = sensitivity.scenario_table(runs)
    assert len(rows) == len(wb.sens_scenarios) == 14
    for got, want in zip(rows, wb.sens_scenarios):
        for key in ("benign", "base", "stress"):
            assert got[key] == pytest.approx(want[key], abs=0.01), f"{want['option']} {key}"


def test_tornado_snapshot(state, wb):
    ids = wb.sensitivity_ids
    assert len(ids) == 23  # the brief says 22; the sheet lists 23
    rows = {r["id"]: r for r in sensitivity.tornado(state, wb.include, ids, wb.units, wb.curves)}
    for want in wb.sens_tornado:
        got = rows[want["id"]]
        assert got["npv_low"] == pytest.approx(want["low"], abs=0.01), f"{want['id']} low"
        assert got["npv_high"] == pytest.approx(want["high"], abs=0.01), f"{want['id']} high"
        assert got["swing"] == pytest.approx(want["swing"], abs=0.01), f"{want['id']} swing"
        assert got["base_npv"] == pytest.approx(want["base"], abs=0.01)


# --- case 2: Stress, bulk margin, 50% exports, recalculated by LibreOffice -----------------

CASE_2 = {"G09": 3, "C11": 1, "G05": 0.5}


@pytest.fixture(scope="module")
def libreoffice_copies(tmp_path_factory):
    """Recalculate two copies of the workbook headless: unchanged, and with CASE_2 typed in."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        pytest.skip("LibreOffice (soffice) is not installed: second parity case skipped")
    work = tmp_path_factory.mktemp("libreoffice")
    src, out = work / "in", work / "out"
    src.mkdir()
    book = openpyxl.load_workbook(DEFAULT_WORKBOOK)  # saving with openpyxl drops cached values
    book.save(src / "base.xlsx")
    ws = book["Inputs"]
    rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1)}
    for iid, value in CASE_2.items():
        ws.cell(rows[iid], 6).value = value  # Base column; the case stays "B"
    book.save(src / "case2.xlsx")
    cmd = [soffice, f"-env:UserInstallation=file://{work / 'profile'}", "--headless", "--norestore",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out),
           str(src / "base.xlsx"), str(src / "case2.xlsx")]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if not (out / "base.xlsx").exists() or not (out / "case2.xlsx").exists():
        pytest.fail(f"LibreOffice did not produce the recalculated files:\n{proc.stdout}\n{proc.stderr}")
    return {name: openpyxl.load_workbook(out / f"{name}.xlsx", data_only=True) for name in ("base", "case2")}


def test_libreoffice_reproduces_excel_for_the_base_case(base, state, wb, libreoffice_copies):
    """Guards the reference: LibreOffice's recalculation equals Excel's cached values."""
    count, problems = _mismatches(base, state, wb, libreoffice_copies["base"])
    assert not problems, f"{len(problems)} of {count} cells differ:\n" + "\n".join(problems[:25])


def test_second_case_matches_libreoffice(state, wb, libreoffice_copies):
    for iid, value in CASE_2.items():
        state.base[iid] = value
    values = state.values()
    assert (values["H01"], values["H08"], values["H06"]) == (0.5, 1100, 1.0)  # hazard rows moved to High
    result = engine.run_model(values, wb.include, wb.units, wb.curves)
    count, problems = _mismatches(result, state, wb, libreoffice_copies["case2"])
    assert not problems, f"{len(problems)} of {count} cells differ:\n" + "\n".join(problems[:25])
    assert count > 6000

    # The same case set the way the app does it (custom values) gives the same run.
    app_state = InputState.from_specs(wb.inputs)
    app_state.custom.update(CASE_2)
    assert app_state.values() == values


# --- speed ---------------------------------------------------------------------------------

def test_engine_runs_well_under_100_ms(state, wb):
    values = state.values()
    engine.run_model(values, wb.include, wb.units, wb.curves)
    start = time.perf_counter()
    for _ in range(50):
        engine.run_model(values, wb.include, wb.units, wb.curves)
    assert (time.perf_counter() - start) / 50 < 0.02
